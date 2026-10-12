from pathlib import Path
import io
from docx import Document
from docx.shared import Inches
from .storage import atomic_write


def export(document: dict, image: Path, target: Path) -> None:
    result = Document()
    result.add_heading(document["relative_path"], level=1)
    result.add_paragraph(f"Provider: {document['provider']} | Model: {document['model']}")
    result.add_paragraph("Complete provider transcript. Engine disagreement alone does not indicate unreadability.")
    for line in document["text"].splitlines():
        result.add_paragraph(line)
    uncertain = [word for line in document["lines"] for word in line.get("words", [])
                 if word.get("confidence") is not None and word["confidence"] < .8]
    if uncertain:
        result.add_heading("Provider confidence below 0.8 (review only)", level=2)
        for word in uncertain:
            result.add_paragraph(f"{word['text']}: {word['confidence']}")
    result.add_heading("OCR working image reference", level=2)
    result.add_picture(str(image), width=Inches(6))
    buffer = io.BytesIO()
    result.save(buffer)
    atomic_write(target, buffer.getvalue())
